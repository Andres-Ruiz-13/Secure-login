#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generate_report.py — Generador de Reportes en Word (.docx) para INBOT SYSTEM / Secure Login
Genera reportes ejecutivos con formato profesional de:
1. Inventario Completo
2. Ventas Diarias / Rango de Fechas
3. Compras por Mes y Año

Uso:
  python generate_report.py --tipo inventario --json datos.json --output reporte_inventario.docx
  python generate_report.py --tipo ventas --json datos.json --output reporte_ventas.docx
  python generate_report.py --tipo compras --json datos.json --output reporte_compras.docx
"""

import sys
import os
import json
import argparse
from datetime import datetime
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Colores de la identidad corporativa
COLOR_PRIMARY_HEX = "0284C7"       # Azul primario
COLOR_DARK_HEX = "090D16"          # Fondo encabezados
COLOR_LIGHT_BG_HEX = "F1F5F9"      # Fondo filas alternas
COLOR_DANGER_HEX = "DC2626"        # Stock bajo
COLOR_SUCCESS_HEX = "16A34A"       # OK / Positivo

def set_cell_background(cell, fill_hex):
    """Aplica color de fondo hexadecimal a una celda de tabla."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Aplica padding interno a una celda (en twips, 1/20 de punto)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tc_pr.append(tcMar)

def add_header(doc, title, subtitle):
    """Agrega un encabezado estilizado al documento."""
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run(title)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(2, 132, 199)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run(f"INBOT SYSTEM & CMS MODULAR | {subtitle}")
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_after = Pt(16)
    run_meta = meta_p.add_run(f"Fecha de emisión: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    run_meta.font.name = 'Arial'
    run_meta.font.size = Pt(9)
    run_meta.font.italic = True
    run_meta.font.color.rgb = RGBColor(148, 163, 184)

def format_table_headers(table, headers):
    """Estiliza la fila de encabezados de una tabla."""
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], COLOR_DARK_HEX)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Arial'
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

def generate_inventory_report(doc, data):
    """Genera reporte completo de inventario."""
    add_header(doc, "REPORTE COMPLETO DE INVENTARIO", "Estado General de Almacén y Existencias")

    # Resumen Ejecutivo
    summary = data.get('resumen', {})
    doc.add_heading("1. Resumen Ejecutivo de Inventario", level=2)
    p_sum = doc.add_paragraph()
    p_sum.add_run(f"• Total Referencias: {summary.get('total_referencias', len(data.get('productos', [])))}\n")
    p_sum.add_run(f"• Total Unidades en Stock: {summary.get('total_unidades', 0):,}\n")
    p_sum.add_run(f"• Valoración al Costo (Compra): ${summary.get('valor_costo_inventario', 0.0):,.2f}\n")
    p_sum.add_run(f"• Valoración Comercial (Venta): ${summary.get('valor_comercial_inventario', 0.0):,.2f}\n")
    p_sum.add_run(f"• Alertas de Stock Bajo: {summary.get('productos_alerta_stock_bajo', 0)}")

    doc.add_heading("2. Detalle de Productos en Existencia", level=2)
    productos = data.get('productos', [])
    headers = ["Código", "Producto", "Categoría", "Stock", "Mínimo", "Precio Compra", "Precio Venta", "Estado"]
    
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_headers(table, headers)

    for idx, p in enumerate(productos):
        row_cells = table.add_row().cells
        stock = int(p.get('stock_actual', 0))
        minimo = int(p.get('stock_minimo', 0))
        es_bajo = stock <= minimo

        row_cells[0].text = str(p.get('codigo_barras', 'N/A'))
        row_cells[1].text = str(p.get('nombre', 'Sin nombre'))
        row_cells[2].text = str(p.get('categoria', 'General'))
        row_cells[3].text = str(stock)
        row_cells[4].text = str(minimo)
        row_cells[5].text = f"${float(p.get('precio_compra', 0)):,.2f}"
        row_cells[6].text = f"${float(p.get('precio_venta', 0)):,.2f}"
        row_cells[7].text = "STOCK BAJO" if es_bajo else "OK"

        bg_color = COLOR_LIGHT_BG_HEX if idx % 2 == 1 else "FFFFFF"
        for i, cell in enumerate(row_cells):
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p_elem = cell.paragraphs[0]
            if i in [0, 3, 4, 5, 6, 7]:
                p_elem.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i in [3, 4, 5, 6] else WD_ALIGN_PARAGRAPH.CENTER
            for run in p_elem.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(8.5)
                if i == 7 and es_bajo:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(220, 38, 38)
                elif i == 7:
                    run.font.color.rgb = RGBColor(22, 163, 74)

def generate_sales_report(doc, data):
    """Genera reporte diario o periódico de ventas."""
    fecha_filtro = data.get('fecha', datetime.now().strftime('%Y-%m-%d'))
    add_header(doc, "REPORTE DE VENTAS", f"Facturación y Transacciones — Fecha: {fecha_filtro}")

    ventas = data.get('ventas', [])
    total_recaudado = sum(float(v.get('total', 0)) for v in ventas)
    doc.add_heading("1. Balance del Período", level=2)
    p_sum = doc.add_paragraph()
    p_sum.add_run(f"• Total Facturas Emitidas: {len(ventas)}\n")
    p_sum.add_run(f"• Total Facturado: ${total_recaudado:,.2f}\n")
    if ventas:
        p_sum.add_run(f"• Ticket Promedio: ${total_recaudado / len(ventas):,.2f}")

    doc.add_heading("2. Relación de Facturas de Venta", level=2)
    headers = ["Factura Nº", "Fecha / Hora", "Cliente", "Método", "Vendedor", "Total ($)"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_headers(table, headers)

    for idx, v in enumerate(ventas):
        row_cells = table.add_row().cells
        row_cells[0].text = str(v.get('numero_factura', 'FAC-000'))
        row_cells[1].text = str(v.get('fecha', 'N/A'))
        row_cells[2].text = str(v.get('cliente_nombre', 'Consumidor Final'))
        row_cells[3].text = str(v.get('metodo_pago', 'Efectivo')).upper()
        row_cells[4].text = str(v.get('vendedor', 'Cajero'))
        row_cells[5].text = f"${float(v.get('total', 0)):,.2f}"

        bg_color = COLOR_LIGHT_BG_HEX if idx % 2 == 1 else "FFFFFF"
        for i, cell in enumerate(row_cells):
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p_elem = cell.paragraphs[0]
            if i in [0, 1, 3]:
                p_elem.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i == 5:
                p_elem.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in p_elem.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(8.5)
                if i in [0, 5]:
                    run.font.bold = True

def generate_purchases_report(doc, data):
    """Genera reporte de compras por mes y año."""
    mes = data.get('mes', datetime.now().strftime('%m'))
    anio = data.get('anio', datetime.now().strftime('%Y'))
    add_header(doc, "REPORTE DE COMPRAS A PROVEEDORES", f"Abastecimiento — Período: {mes}/{anio}")

    compras = data.get('compras', [])
    total_compras = sum(float(c.get('total', 0)) for c in compras)
    doc.add_heading("1. Resumen de Adquisiciones", level=2)
    p_sum = doc.add_paragraph()
    p_sum.add_run(f"• Total Facturas de Compra: {len(compras)}\n")
    p_sum.add_run(f"• Total Inversión en Compras: ${total_compras:,.2f}\n")

    doc.add_heading("2. Relación de Facturas de Compra", level=2)
    headers = ["Factura Nº", "Fecha", "Proveedor", "NIT", "Registrado Por", "Total ($)"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    format_table_headers(table, headers)

    for idx, c in enumerate(compras):
        row_cells = table.add_row().cells
        row_cells[0].text = str(c.get('numero_factura', 'COM-000'))
        row_cells[1].text = str(c.get('fecha', 'N/A'))
        row_cells[2].text = str(c.get('proveedor', 'N/A'))
        row_cells[3].text = str(c.get('nit', 'N/A'))
        row_cells[4].text = str(c.get('usuario', 'Bodega'))
        row_cells[5].text = f"${float(c.get('total', 0)):,.2f}"

        bg_color = COLOR_LIGHT_BG_HEX if idx % 2 == 1 else "FFFFFF"
        for i, cell in enumerate(row_cells):
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p_elem = cell.paragraphs[0]
            if i in [0, 1, 3]:
                p_elem.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i == 5:
                p_elem.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in p_elem.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(8.5)
                if i in [0, 5]:
                    run.font.bold = True

def main():
    parser = argparse.ArgumentParser(description="Generador de reportes Word (.docx) para INBOT SYSTEM")
    parser.add_argument("--tipo", choices=["inventario", "ventas", "compras"], required=True, help="Tipo de reporte")
    parser.add_argument("--json", required=True, help="Ruta al archivo JSON de datos")
    parser.add_argument("--output", required=True, help="Ruta del archivo .docx de salida")
    args = parser.parse_args()

    if not os.path.exists(args.json):
        print(f"Error: No existe el archivo JSON {args.json}", file=sys.stderr)
        sys.exit(1)

    with open(args.json, 'r', encoding='utf-8') as f:
        data = json.load(f)

    doc = Document()

    # Ajuste de márgenes a 2 cm (aprox 0.78 pulgadas)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    if args.tipo == "inventario":
        generate_inventory_report(doc, data)
    elif args.tipo == "ventas":
        generate_sales_report(doc, data)
    elif args.tipo == "compras":
        generate_purchases_report(doc, data)

    # Crear directorio padre si no existe
    out_dir = os.path.dirname(os.path.abspath(args.output))
    if out_dir and not os.path.exists(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    doc.save(args.output)
    print(f"Reporte generado exitosamente en: {args.output}")

if __name__ == "__main__":
    main()

